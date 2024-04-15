import re
import os
import docx
from openpyxl import Workbook
from PyPDF2 import PdfReader

def extract_info_from_cv(text):
    email_pattern = r'[\w\.-]+@[\w\.-]+'
    phone_pattern = r'\b(?:\d{3}[-\.\s]?)?\d{3}[-\.\s]?\d{4}\b'

    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)

    # Remove duplicates
    emails = list(set(emails))
    phones = list(set(phones))

    # Assuming overall text is everything except emails and phones
    overall_text = re.sub(email_pattern, '', text)
    overall_text = re.sub(phone_pattern, '', overall_text)

    return emails, phones, overall_text
def save_to_excel(emails, phones, overall_text, output_file="cv_data_pdf.xlsx"):
    # Create a new Excel workbook
    wb = Workbook()
    ws = wb.active

    # Add headers to the worksheet
    ws.append(["Email", "Phone", "Overall Text"])

    # Add data to the worksheet
    for email in emails:
        ws.append([email, "", ""])
    for phone in phones:
        ws.append(["", phone, ""])
    ws.append(["", "", overall_text])

    # Save the workbook to a file
    wb.save(output_file)


newpath = input("Enter the path of your file: ")

# Check if the file exists
if not os.path.exists(newpath):
    print("File not found at the specified path.")
else:
    forpath = os.path.splitext(newpath)
    forextract = forpath[1]

    if forextract == ".pdf":
        def extract_text_from_pdf(pdf_path):
            with open(pdf_path, 'rb') as pdf_file:
                reader = PdfReader(pdf_file)
                text = ''
                for page_num in range(len(reader.pages)):
                    text += reader.pages[page_num].extract_text()
            return text

        try:
            emails, phones, overall_text = extract_info_from_cv(extract_text_from_pdf(newpath))
            save_to_excel(emails, phones, overall_text, "cv_data_pdf.xlsx")
        except Exception as e:
            print("Error extracting text from PDF:", e)

    elif forextract == ".docx" or ".doc":
        def extract_text_from_docx(docx_path):
            doc = docx.Document(docx_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text 
            return text

        try:
            emails, phones, overall_text = extract_info_from_cv(extract_text_from_docx(newpath))
            save_to_excel(emails, phones, overall_text)
        except Exception as e:
            print("Error extracting text from DOCX:", e)
